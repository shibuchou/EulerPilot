#!/usr/bin/env python3
"""Render the Chinese deliverable documents (docs/*.md) into Word .docx.

Markdown is the source of truth; table header shading, page numbers and CJK fonts
cannot live in markdown, so they are applied here. Pandoc supplies the structure,
a styles patch supplies the blue theme, and a document pass paints header rows
because Word does not always honour table-style conditional shading.
"""
from __future__ import annotations

import argparse
import hashlib
import os
import re
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path

REPO = Path(__file__).resolve().parents[1]
DOC_DIR = REPO / "docs"
TEMPLATE = DOC_DIR / "templates" / "reference-blue.docx"

ACCENT = "1F4E79"
ACCENT_SOFT = "2E6DA4"
HEADER_FILL = "1F4E79"
BAND_FILL = "EDF3F9"
CODE_FILL = "F4F6F8"
GREY = "595959"
HEADING_FONT = "微软雅黑"
BODY_FONT = "宋体"
LATIN_BODY = "Calibri"
CODE_FONT = "Consolas"

DOCS = {
    "项目说明书": ("EulerPilot-v2项目说明书", "v1.1"),
    "设计文档": ("EulerPilot-v2设计文档", "v1.1"),
    "用户手册": ("EulerPilot-v2用户手册", "v1.1"),
    "部署文档": ("EulerPilot-v2部署文档", "v1.1"),
    "测试文档": ("EulerPilot-v2测试文档", "v1.1"),
    "REPORT": ("EulerPilot-v2参赛文档", "v1.1"),
}

TBL_BORDERS = (
    '<w:tblBorders>'
    '<w:top w:val="single" w:sz="8" w:space="0" w:color="9BB7D4"/>'
    '<w:left w:val="single" w:sz="8" w:space="0" w:color="9BB7D4"/>'
    '<w:bottom w:val="single" w:sz="8" w:space="0" w:color="9BB7D4"/>'
    '<w:right w:val="single" w:sz="8" w:space="0" w:color="9BB7D4"/>'
    '<w:insideH w:val="single" w:sz="4" w:space="0" w:color="C5D5E6"/>'
    '<w:insideV w:val="single" w:sz="4" w:space="0" w:color="C5D5E6"/>'
    "</w:tblBorders>"
)


def heading_rpr(size: int, color: str) -> str:
    return (
        "<w:rPr>"
        f'<w:rFonts w:ascii="{HEADING_FONT}" w:hAnsi="{HEADING_FONT}" '
        f'w:eastAsia="{HEADING_FONT}" w:cs="{HEADING_FONT}"/>'
        "<w:b /><w:color w:val=\"" + color + "\" />"
        f'<w:sz w:val="{size}" /><w:szCs w:val="{size}" />'
        "</w:rPr>"
    )


def patch_styles(styles: str) -> str:
    styles = re.sub(
        r"<w:rFonts w:asciiTheme=\"minorHAnsi\"[^/]*/>",
        f'<w:rFonts w:ascii="{LATIN_BODY}" w:hAnsi="{LATIN_BODY}" '
        f'w:eastAsia="{BODY_FONT}" w:cs="{LATIN_BODY}" />',
        styles,
        count=1,
    )
    styles = re.sub(
        r"<w:sz w:val=\"22\" /><w:szCs w:val=\"22\" />",
        '<w:sz w:val="21" /><w:szCs w:val="21" />',
        styles,
        count=1,
    )
    sizes = {1: (34, ACCENT), 2: (28, ACCENT), 3: (24, ACCENT_SOFT), 4: (22, ACCENT_SOFT),
             5: (21, GREY), 6: (21, GREY)}
    for level, (size, color) in sizes.items():
        sid = f"Heading{level}"
        block = re.search(r'<w:style [^>]*w:styleId="%s">.*?</w:style>' % sid, styles, re.S)
        if not block:
            continue
        new = re.sub(r"<w:rPr>.*?</w:rPr>", heading_rpr(size, color), block.group(0), count=1, flags=re.S)
        new = re.sub(r"<w:spacing([^/]*)/>", r'<w:spacing\1 w:before="' + str(320 - 40 * level) +
                     '" w:after="120" w:line="264" w:lineRule="auto"/>', new, count=1)
        styles = styles.replace(block.group(0), new)
    table = re.search(r'<w:style w:type="table"[^>]*w:styleId="Table">.*?</w:style>', styles, re.S)
    if table:
        replacement = (
            '<w:style w:type="table" w:default="1" w:styleId="Table">'
            '<w:name w:val="Table" /><w:basedOn w:val="TableNormal" />'
            "<w:tblPr>"
            '<w:tblInd w:w="0" w:type="dxa" />' + TBL_BORDERS +
            '<w:tblCellMar><w:top w:w="60" w:type="dxa" />'
            '<w:left w:w="108" w:type="dxa" /><w:bottom w:w="60" w:type="dxa" />'
            '<w:right w:w="108" w:type="dxa" /></w:tblCellMar>'
            "</w:tblPr>"
            '<w:tblStylePr w:type="firstRow"><w:rPr><w:b /><w:color w:val="FFFFFF" />'
            f'<w:sz w:val="20" /><w:szCs w:val="20" /></w:rPr>'
            '<w:tcPr><w:tcBorders>'
            '<w:top w:val="single" w:sz="8" w:space="0" w:color="9BB7D4"/>'
            '<w:bottom w:val="single" w:sz="8" w:space="0" w:color="9BB7D4"/>'
            "</w:tcBorders>"
            f'<w:shd w:val="clear" w:color="auto" w:fill="{HEADER_FILL}" />'
            '<w:vAlign w:val="center" /></w:tcPr></w:tblStylePr>'
            '<w:tblStylePr w:type="lastRow"><w:tcPr><w:tcBorders>'
            '<w:bottom w:val="single" w:sz="8" w:space="0" w:color="9BB7D4"/>'
            "</w:tcBorders></w:tcPr></w:tblStylePr>"
            '<w:tblStylePr w:type="band1Horz"><w:tcPr>'
            f'<w:shd w:val="clear" w:color="auto" w:fill="{BAND_FILL}" />'
            "</w:tcPr></w:tblStylePr>"
            "</w:style>"
        )
        styles = styles.replace(table.group(0), replacement)
    for sid in ("SourceCode", "VerbatimChar"):
        block = re.search(r'<w:style [^>]*w:styleId="%s">.*?</w:style>' % sid, styles, re.S)
        if not block:
            continue
        body = block.group(0)
        fonts = f'<w:rFonts w:ascii="{CODE_FONT}" w:hAnsi="{CODE_FONT}" w:eastAsia="{BODY_FONT}" />'
        if "<w:rPr>" in body:
            new = body.replace("<w:rPr>", "<w:rPr>" + fonts + f'<w:sz w:val="18" /><w:szCs w:val="18" />', 1)
        else:
            new = body[: -len("</w:style>")] + "<w:rPr>" + fonts + f'<w:sz w:val="18" /><w:szCs w:val="18" /></w:rPr></w:style>'
        styles = styles.replace(body, new)
    return styles


def patch_code_paragraphs(styles: str) -> str:
    """Give code blocks a light panel, matching the reference manual's shaded listings."""
    block = re.search(r'(<w:style w:type="paragraph"[^>]*w:styleId="SourceCode">)(.*?)(</w:style>)', styles, re.S)
    if not block:
        return styles
    head, body, tail = block.group(1), block.group(2), block.group(3)
    panel = (
        '<w:pPr><w:keepNext />'
        '<w:pBdr><w:left w:val="single" w:sz="18" w:space="4" w:color="' + ACCENT + '" /></w:pBdr>'
        f'<w:shd w:val="clear" w:color="auto" w:fill="{CODE_FILL}" />'
        '<w:spacing w:before="60" w:after="60" w:line="240" w:lineRule="auto" />'
        '<w:ind w:left="180" /></w:pPr>'
    )
    body = re.sub(r"<w:pPr>.*?</w:pPr>", "", body, count=1, flags=re.S)
    body = re.sub(r"<w:rPr>", panel + "<w:rPr>", body, count=1)
    if "<w:rPr>" not in body:
        body += panel
    return styles.replace(block.group(0), head + body + tail)


def build_template() -> None:
    ref = subprocess.run([
        "pandoc", "--print-default-data-file", "reference.docx"
    ], check=True, capture_output=True).stdout
    with tempfile.TemporaryDirectory() as td:
        seed = Path(td) / "ref.docx"
        seed.write_bytes(ref)
        work = Path(td) / "x"
        with zipfile.ZipFile(seed) as z:
            z.extractall(work)
        styles_path = work / "word" / "styles.xml"
        styles = styles_path.read_text(encoding="utf-8")
        styles = patch_styles(styles)
        styles = patch_code_paragraphs(styles)
        styles_path.write_text(styles, encoding="utf-8")
        TEMPLATE.parent.mkdir(parents=True, exist_ok=True)
        if TEMPLATE.exists():
            TEMPLATE.unlink()
        with zipfile.ZipFile(TEMPLATE, "w", zipfile.ZIP_DEFLATED) as z:
            for path in sorted(work.rglob("*")):
                if path.is_file():
                    z.write(path, path.relative_to(work).as_posix())


def paint_header_rows(document: str) -> str:
    def fix_table(match: re.Match) -> str:
        xml = match.group(0)
        xml = re.sub(r'<w:tblW w:type="auto" w:w="0" />', '<w:tblW w:type="pct" w:w="5000" />', xml)
        if "<w:tblBorders>" not in xml:
            xml = xml.replace("<w:tblLook", TBL_BORDERS + "<w:tblLook", 1)
        rows = re.findall(r"<w:tr>.*?</w:tr>", xml, re.S)
        for row in rows:
            if "<w:tblHeader" not in row:
                continue
            new_row = re.sub(
                r"<w:tcPr\s*/>",
                '<w:tcPr><w:shd w:val="clear" w:color="auto" w:fill="' + HEADER_FILL + '" />'
                '<w:vAlign w:val="center" /></w:tcPr>',
                row,
            )
            new_row = re.sub(
                r"<w:r><w:t( xml:space=\"preserve\")?>",
                lambda r: '<w:r><w:rPr><w:b /><w:color w:val="FFFFFF" /><w:sz w:val="20" />'
                f'<w:szCs w:val="20" /></w:rPr><w:t{r.group(1) or ""}>',
                new_row,
            )
            xml = xml.replace(row, new_row, 1)
        return xml

    return re.sub(r"<w:tbl>.*?</w:tbl>", fix_table, document, flags=re.S)


def page_field(instr: str) -> str:
    return (
        '<w:fldSimple xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
        f'w:instr="{instr}"><w:r><w:t>1</w:t></w:r></w:fldSimple>'
    )


def polish(path: Path, title: str) -> None:
    with zipfile.ZipFile(path) as z:
        names = z.namelist()
        parts = {n: z.read(n) for n in names}
    document = parts["word/document.xml"].decode("utf-8")
    document = paint_header_rows(document)
    parts["word/document.xml"] = document.encode("utf-8")
    styles = parts["word/styles.xml"].decode("utf-8")
    parts["word/styles.xml"] = patch_code_paragraphs(styles).encode("utf-8")
    path.unlink()
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        for name in names:
            z.writestr(name, parts[name])

    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import parse_xml
    from docx.shared import Pt, RGBColor, Cm

    doc = Document(str(path))
    for section in doc.sections:
        section.page_width, section.page_height = Cm(21.0), Cm(29.7)
        section.top_margin = section.bottom_margin = Cm(2.2)
        section.left_margin = section.right_margin = Cm(2.4)
        header = section.header.paragraphs[0]
        header.text = ""
        run = header.add_run(title)
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(0x59, 0x59, 0x59)
        run.font.name = HEADING_FONT
        header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        footer = section.footer.paragraphs[0]
        footer.text = ""
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        prefix = footer.add_run("第 ")
        prefix.font.size = Pt(9)
        footer._p.append(parse_xml(page_field("PAGE")))
        middle = footer.add_run(" 页 / 共 ")
        middle.font.size = Pt(9)
        footer._p.append(parse_xml(page_field("NUMPAGES")))
        tail = footer.add_run(" 页")
        tail.font.size = Pt(9)
    props = doc.core_properties
    props.title = title
    props.author = "三驱兄弟（西安邮电大学）"
    props.category = "2026 年全国大学生计算机系统能力大赛操作系统设计赛"
    for paragraph in doc.paragraphs:
        if paragraph._p.xpath('.//w:drawing'):
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.keep_with_next = True
            paragraph.paragraph_format.space_before = Pt(8)
            paragraph.paragraph_format.space_after = Pt(4)
        elif re.match(r'^图 \d+\s', paragraph.text):
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.keep_with_next = False
            paragraph.paragraph_format.space_after = Pt(10)
            for run in paragraph.runs:
                run.font.size = Pt(9)
                run.font.color.rgb = RGBColor.from_string(GREY)
    # Pandoc represents a whole listing as one paragraph. Keeping it with the
    # following heading/figure creates long empty page tails; allow natural
    # listing continuation and preserve the image/caption pairing above.
    for paragraph in doc.paragraphs:
        if paragraph.style and paragraph.style.style_id == 'SourceCode':
            paragraph.paragraph_format.keep_with_next = False
    # Revision rows should remain readable and keep a short date/author column.
    for table in doc.tables:
        headers = [cell.text.strip() for cell in table.rows[0].cells]
        if headers == ['版本', '日期', '内容', '作者']:
            table.autofit = False
            for index, row in enumerate(table.rows):
                for cell, width in zip(row.cells, [2.4,2.5,8.8,2.5]):
                    cell.width = Cm(width)
                    for paragraph in cell.paragraphs:
                        paragraph.paragraph_format.keep_with_next = index < len(table.rows) - 1
                props = row._tr.get_or_add_trPr()
                props.append(parse_xml('<w:cantSplit xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"/>'))
    # Word requires a paragraph after the final table. A full-size empty one can
    # generate a blank trailing page when the last row ends at the page bottom.
    body_items = [element for element in doc._element.body
                  if element.tag.endswith('}p') or element.tag.endswith('}tbl')]
    if body_items and body_items[-1].tag.endswith('}tbl'):
        doc.add_paragraph()
    if doc.paragraphs and not doc.paragraphs[-1].text.strip():
        last = doc.paragraphs[-1]
        last.paragraph_format.space_before = Pt(0)
        last.paragraph_format.space_after = Pt(0)
        last.paragraph_format.line_spacing = Pt(1)
        last.paragraph_format.keep_with_next = False
        last._p.get_or_add_pPr().append(parse_xml(
            '<w:rPr xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
            '<w:sz w:val="2"/><w:szCs w:val="2"/></w:rPr>'))
    doc.save(str(path))


def to_pdf(docx_path: Path) -> Path:
    """Export through Word itself; the docx must survive that open, which doubles as a validity check."""
    try:
        import pythoncom
        from win32com import client
    except ImportError:
        if sys.platform != "win32":
            raise RuntimeError("PDF export requires Microsoft Word on Windows")
        subprocess.run([
            "powershell.exe", "-NoProfile", "-NonInteractive", "-ExecutionPolicy", "Bypass",
            "-File", str(REPO / "scripts/export_docx_pdf.ps1"),
            "-DocxPath", str(docx_path.resolve()),
        ], check=True, timeout=180)
        pdf = docx_path.with_suffix('.pdf')
        if not pdf.is_file() or pdf.stat().st_size == 0:
            raise RuntimeError(f"Word did not produce {pdf}")
        return pdf

    pdf = docx_path.with_suffix(".pdf")
    pythoncom.CoInitialize()
    app = client.DispatchEx("Word.Application")
    app.Visible = False
    app.DisplayAlerts = 0
    try:
        doc = app.Documents.Open(str(docx_path), ReadOnly=True, Visible=False)
        doc.ExportAsFixedFormat(str(pdf), 17)
        doc.Close(False)
    finally:
        app.Quit()
    return pdf


def prepare_figures(source: str, md: Path, png_dir: Path):
    """Use high-resolution previews for Pandoc; embed native SVG after polishing."""
    assets = {}
    def replace(match):
        alt, target = match.group(1), match.group(2)
        svg = (md.parent / target).resolve()
        if not svg.is_file():
            raise FileNotFoundError(svg)
        png = png_dir / (svg.stem + '.png')
        if not png.is_file():
            raise FileNotFoundError(f"Missing figure preview {png}; run node scripts/render_diagrams.cjs first")
        assets[hashlib.sha256(png.read_bytes()).hexdigest()] = svg
        return f'![{alt}](<{png.resolve().as_posix()}>){{width=6.35in}}'
    source = re.sub(r'!\[([^\]]*)\]\(([^)]+\.svg)\)', replace, source)
    return source, assets


def embed_native_svg(path: Path, assets: dict) -> None:
    """Add Office SVG blips while retaining PNG fallbacks for older readers."""
    if not assets:
        return
    from lxml import etree as ET
    with zipfile.ZipFile(path) as archive:
        parts = {n: archive.read(n) for n in archive.namelist()}
    ns = {
        'a': 'http://schemas.openxmlformats.org/drawingml/2006/main',
        'r': 'http://schemas.openxmlformats.org/officeDocument/2006/relationships',
        'asvg': 'http://schemas.microsoft.com/office/drawing/2016/SVG/main',
        'p': 'http://schemas.openxmlformats.org/package/2006/relationships',
        'ct': 'http://schemas.openxmlformats.org/package/2006/content-types',
    }
    document = ET.fromstring(parts['word/document.xml'])
    relationships = ET.fromstring(parts['word/_rels/document.xml.rels'])
    types = ET.fromstring(parts['[Content_Types].xml'])
    mapping = {}
    for rel in list(relationships):
        target = rel.get('Target', '')
        part = 'word/' + target
        if not target.endswith('.png') or part not in parts:
            continue
        svg = assets.get(hashlib.sha256(parts[part]).hexdigest())
        if svg is None:
            continue
        svg_target = 'media/' + svg.name
        parts['word/' + svg_target] = svg.read_bytes()
        rid = 'rIdSvg' + str(len(mapping) + 1)
        if relationships.xpath('*[@Id=$rid]', rid=rid):
            raise ValueError(f'duplicate relationship id: {rid}')
        ET.SubElement(relationships, '{'+ns['p']+'}Relationship',
                      Id=rid, Type=ns['r']+'/image', Target=svg_target)
        mapping[rel.get('Id')] = rid
    if len(mapping) < len(assets):
        raise ValueError(f'Expected {len(assets)} native SVG assets, found {len(mapping)}')
    for blip in document.findall('.//a:blip', ns):
        svg_id = mapping.get(blip.get('{'+ns['r']+'}embed'))
        if svg_id is None:
            continue
        ext_list = blip.find('a:extLst', ns)
        if ext_list is None:
            ext_list = ET.SubElement(blip, '{'+ns['a']+'}extLst')
        ext = ET.SubElement(ext_list, '{'+ns['a']+'}ext', uri='{96DAC541-7B7A-43D3-8B79-37D633B846F1}')
        ET.SubElement(ext, '{'+ns['asvg']+'}svgBlip',
                      {'{'+ns['r']+'}embed': svg_id}, nsmap={'asvg':ns['asvg']})
    if not types.xpath('*[@Extension="svg"]'):
        ET.SubElement(types, '{'+ns['ct']+'}Default', Extension='svg', ContentType='image/svg+xml')
    for name, element in [('word/document.xml',document),('word/_rels/document.xml.rels',relationships),('[Content_Types].xml',types)]:
        parts[name] = ET.tostring(element, encoding='UTF-8', xml_declaration=True, standalone=True)
    with zipfile.ZipFile(path, 'w', zipfile.ZIP_DEFLATED) as archive:
        for name, data in parts.items():
            archive.writestr(name, data)


def convert(name: str, force_template: bool = False, output_dir: Path = DOC_DIR,
            png_dir: Path = REPO / 'build/diagram-qa/png') -> Path:
    stem, version = DOCS[name]
    md = DOC_DIR / f"{name}.md"
    if not md.exists():
        raise SystemExit(f"missing source: {md}")
    if force_template or not TEMPLATE.exists():
        build_template()
    output_dir.mkdir(parents=True, exist_ok=True)
    out = output_dir / f"{stem}-{version}.docx"
    prepared, assets = prepare_figures(md.read_text(encoding='utf-8'), md, png_dir)
    with tempfile.TemporaryDirectory(prefix='ep-docs-') as temp:
        source = Path(temp) / md.name
        source.write_text(prepared, encoding='utf-8')
        subprocess.run([
            "pandoc", str(source), "-f", "markdown+east_asian_line_breaks+raw_html-implicit_figures",
            "--resource-path", os.pathsep.join([str(DOC_DIR), str(REPO)]),
            "-t", "docx", "--reference-doc", str(TEMPLATE), "-o", str(out),
        ], check=True)
    polish(out, f"{stem} {version}")
    embed_native_svg(out, assets)
    return out


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("docs", nargs="*", help="document basenames without .md")
    ap.add_argument("--all", action="store_true")
    ap.add_argument("--rebuild-template", action="store_true")
    ap.add_argument("--pdf", action="store_true", help="also export PDF through Word")
    ap.add_argument("--output-dir", type=Path, default=DOC_DIR, help="stage outputs in this directory")
    ap.add_argument("--figure-png-dir", type=Path, default=REPO / 'build/diagram-qa/png')
    args = ap.parse_args()
    targets = list(DOCS) if args.all or not args.docs else args.docs
    for name in targets:
        if name not in DOCS:
            raise SystemExit(f"unknown document: {name} (known: {', '.join(DOCS)})")
        out = convert(name, force_template=args.rebuild_template,
                      output_dir=args.output_dir.resolve(), png_dir=args.figure_png_dir.resolve())
        print(f"{out}  ({out.stat().st_size // 1024} KiB)")
        if args.pdf:
            pdf = to_pdf(out)
            print(f"{pdf}  ({pdf.stat().st_size // 1024} KiB)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
